#!/usr/bin/env python3
"""Consolida reportes Excel META4 en hojas Espera y Cumplimiento."""

from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
import math
from pathlib import Path
import re
import sys
import unicodedata
from typing import Any


APP_VERSION = "1.2"
EXCEL_EXTENSIONS = {".xls", ".xlsx", ".xlsm"}
OUTPUT_BASENAME = "Registros consolidados"

HEADERS = {
    "Espera": [
        "RIT",
        "TRIBUNAL",
        "RUT",
        "NOMBRE",
        "NACIONALIDAD / PAÍS",
        "TIPO",
        "SEXO",
        "DERIVACIÓN",
        "CURADOR",
        "T ESPERA",
        "FEC. RESOLUCIÓN",
        "REVISIÓN",
        "DURACIÓN",
        "80 BIS",
        "FEC. NACIMIENTO",
        "EDAD",
        "FEC. ACT. F. INDIVIDUAL",
        "FEC. ACT. F. AMBULATORIA",
        "FEC. ACT. F. FAE/FAS",
        "FEC. OÍDO",
        "PROXS. AUDS.",
    ],
    "Cumplimiento": [
        "RIT",
        "TRIBUNAL",
        "RUT",
        "NOMBRE",
        "NACIONALIDAD / PAÍS",
        "TIPO",
        "SEXO",
        "DERIVACIÓN",
        "CURADOR",
        "FEC. RESOLUCIÓN",
        "FEC. INGRESO EFECTIVO",
        "FEC. EGRESO PROYECTADO",
        "DÍAS DE CUMPLIMIENTO",
        "DÍAS PARA EGRESAR",
        "80 BIS",
        "FEC. NACIMIENTO",
        "EDAD",
        "FEC. ACT. F. INDIVIDUAL",
        "FEC. ACT. F. AMBULATORIA",
        "FEC. ACT. F. FAE/FAS",
        "FEC. OÍDO",
    ],
}

DUPLICATE_HEADERS = [
    "UBICACIÓN",
    "RIT",
    "TRIBUNAL",
    "RUT",
    "NOMBRE",
    "NACIONALIDAD / PAÍS",
    "TIPO",
    "SEXO",
    "DERIVACIÓN",
    "CURADOR",
    "T ESPERA",
    "FEC. RESOLUCIÓN",
    "FEC. INGRESO EFECTIVO",
    "FEC. EGRESO PROYECTADO",
    "DÍAS DE CUMPLIMIENTO",
    "DÍAS PARA EGRESAR",
    "REVISIÓN",
    "DURACIÓN",
    "80 BIS",
    "FEC. NACIMIENTO",
    "EDAD",
    "FEC. ACT. F. INDIVIDUAL",
    "FEC. ACT. F. AMBULATORIA",
    "FEC. ACT. F. FAE/FAS",
    "FEC. OÍDO",
    "PROXS. AUDS.",
]


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]


@dataclass
class ExtractedFile:
    source: Path
    status: str
    sheet_name: str
    header_row: int
    records: list[list[Any]]


def normalize(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def import_xlrd():
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "Falta xlrd para leer archivos .xls. Ejecuta Instalar dependencias.bat"
        ) from exc
    return xlrd


def import_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Falta openpyxl. Ejecuta Instalar dependencias.bat"
        ) from exc
    return openpyxl


def read_xls(path: Path) -> list[SheetData]:
    xlrd = import_xlrd()
    book = xlrd.open_workbook(str(path), on_demand=True)
    sheets: list[SheetData] = []
    try:
        for name in book.sheet_names():
            sheet = book.sheet_by_name(name)
            rows: list[list[Any]] = []
            for row_index in range(sheet.nrows):
                row: list[Any] = []
                for col_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, col_index)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    else:
                        value = cell.value
                    row.append(value)
                rows.append(row)
            sheets.append(SheetData(name=name, rows=rows))
    finally:
        book.release_resources()
    return sheets


def read_xlsx(path: Path) -> list[SheetData]:
    openpyxl = import_openpyxl()
    book = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )
    try:
        return [
            SheetData(
                name=sheet.title,
                rows=[list(row) for row in sheet.iter_rows(values_only=True)],
            )
            for sheet in book.worksheets
        ]
    finally:
        book.close()


def read_workbook(path: Path) -> list[SheetData]:
    if path.suffix.lower() == ".xls":
        return read_xls(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise ValueError(f"Extensión no compatible: {path.suffix}")


def find_header_row(rows: list[list[Any]], max_rows: int = 30) -> int | None:
    for row_index, row in enumerate(rows[:max_rows]):
        values = {normalize(value) for value in row if normalize(value)}
        if (
            "RIT" in values
            and "TRIBUNAL" in values
            and any(value.startswith("DERIV") for value in values)
        ):
            return row_index
    return None


def header_map(headers: list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(headers):
        key = normalize(value)
        if key and key not in result:
            result[key] = index
    return result


def get_header_index(mapping: dict[str, int], target: str) -> int | None:
    wanted = normalize(target)
    if wanted in mapping:
        return mapping[wanted]
    if wanted == "DERIVACION":
        return next(
            (index for key, index in mapping.items() if key.startswith("DERIV")),
            None,
        )
    return None


def detect_status(title: Any, headers: list[Any]) -> str:
    title_text = normalize(title)
    title_status: str | None = None
    if "LISTA DE ESPERA" in title_text:
        title_status = "Espera"
    elif "LISTA DE CUMPLIMIENTO" in title_text:
        title_status = "Cumplimiento"

    normalized_headers = {normalize(value) for value in headers}
    header_status: str | None = None
    if "T ESPERA" in normalized_headers:
        header_status = "Espera"
    if (
        "FEC INGRESO EFECTIVO" in normalized_headers
        or "DIAS DE CUMPLIMIENTO" in normalized_headers
    ):
        if header_status and header_status != "Cumplimiento":
            raise ValueError("los encabezados mezclan Espera y Cumplimiento")
        header_status = "Cumplimiento"

    if title_status and header_status and title_status != header_status:
        raise ValueError("el título contradice el estado indicado por los encabezados")
    status = title_status or header_status
    if not status:
        raise ValueError("no fue posible identificar Espera o Cumplimiento por contenido")
    return status


def value_at(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def is_record_row(row: list[Any], rit_index: int) -> bool:
    value = value_at(row, rit_index)
    return value is not None and str(value).strip() != ""


def select_data_sheet(sheets: list[SheetData]) -> tuple[SheetData, int]:
    candidates: list[tuple[int, int, str, SheetData, int]] = []
    for sheet in sheets:
        row_index = find_header_row(sheet.rows)
        if row_index is None:
            continue
        mapping = header_map(sheet.rows[row_index])
        rit_index = get_header_index(mapping, "RIT")
        if rit_index is None:
            continue
        record_count = sum(
            is_record_row(row, rit_index) for row in sheet.rows[row_index + 1 :]
        )
        header_score = len(mapping)
        candidates.append((record_count, header_score, sheet.name, sheet, row_index))

    if not candidates:
        raise ValueError("no se encontró una hoja con RIT, TRIBUNAL y DERIVACIÓN")
    candidates.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    _, _, _, sheet, row_index = candidates[0]
    return sheet, row_index


def extract_file(path: Path) -> ExtractedFile:
    sheets = read_workbook(path)
    sheet, header_row = select_data_sheet(sheets)
    headers = sheet.rows[header_row]
    title = sheet.rows[1][0] if len(sheet.rows) > 1 and sheet.rows[1] else None
    status = detect_status(title, headers)
    mapping = header_map(headers)
    rit_index = get_header_index(mapping, "RIT")
    if rit_index is None:
        raise ValueError("no se encontró la columna RIT")

    missing = [
        header
        for header in HEADERS[status]
        if get_header_index(mapping, header) is None
    ]
    if missing:
        raise ValueError("faltan columnas: " + ", ".join(missing))

    column_indexes = [get_header_index(mapping, header) for header in HEADERS[status]]
    records = [
        [value_at(row, index) for index in column_indexes]
        for row in sheet.rows[header_row + 1 :]
        if is_record_row(row, rit_index)
    ]
    return ExtractedFile(
        source=path,
        status=status,
        sheet_name=sheet.name,
        header_row=header_row + 1,
        records=records,
    )


def is_prior_output(path: Path) -> bool:
    stem = normalize(path.stem)
    if stem == normalize(OUTPUT_BASENAME) or re.fullmatch(
        rf"{re.escape(normalize(OUTPUT_BASENAME))} [0-9]+", stem
    ):
        return True
    for part in path.parts:
        normalized = normalize(part)
        if normalized == "ARCHIVOS ETIQUETADOS" or re.fullmatch(
            r"ARCHIVOS ETIQUETADOS [0-9]+", normalized
        ):
            return True
    return False


def find_excels(folder: Path) -> list[Path]:
    return sorted(
        path
        for path in folder.rglob("*")
        if path.is_file()
        and path.suffix.lower() in EXCEL_EXTENSIONS
        and not path.name.startswith("~$")
        and not is_prior_output(path.relative_to(folder))
    )


def next_output_path(folder: Path) -> Path:
    first = folder / f"{OUTPUT_BASENAME}.xlsx"
    if not first.exists():
        return first
    number = 2
    while True:
        candidate = folder / f"{OUTPUT_BASENAME} {number}.xlsx"
        if not candidate.exists():
            return candidate
        number += 1


def safe_excel_value(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def normalize_rut(value: Any) -> str:
    return re.sub(r"[^0-9K]+", "", str(value or "").upper())


def find_duplicate_records(
    records: dict[str, list[list[Any]]],
) -> tuple[list[list[Any]], int]:
    groups: dict[tuple[str, str], list[tuple[str, list[Any]]]] = defaultdict(list)

    for status in ("Espera", "Cumplimiento"):
        rut_index = HEADERS[status].index("RUT")
        name_index = HEADERS[status].index("NOMBRE")
        for row in records[status]:
            key = (normalize_rut(value_at(row, rut_index)), normalize(value_at(row, name_index)))
            if key[0] and key[1]:
                groups[key].append((status, row))

    duplicate_groups = {
        key: occurrences for key, occurrences in groups.items() if len(occurrences) > 1
    }
    duplicate_rows: list[list[Any]] = []

    for key in sorted(duplicate_groups):
        occurrences = sorted(
            duplicate_groups[key],
            key=lambda item: (item[0] != "Espera", normalize(value_at(item[1], 0))),
        )
        for status, row in occurrences:
            source = {
                normalize(header): value_at(row, index)
                for index, header in enumerate(HEADERS[status])
            }
            duplicate_rows.append(
                [status]
                + [source.get(normalize(header)) for header in DUPLICATE_HEADERS[1:]]
            )

    return duplicate_rows, len(duplicate_groups)


def create_output(
    output_path: Path,
    records: dict[str, list[list[Any]]],
    duplicate_rows: list[list[Any]],
) -> None:
    openpyxl = import_openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    default_sheet = book.active
    book.remove(default_sheet)

    normal_widths = [
        16, 28, 15, 34, 22, 15, 12, 42, 42, 16, 18, 16, 18, 18, 15, 18,
        12, 22, 24, 22, 24,
    ]
    duplicate_widths = [
        18, 16, 28, 15, 34, 22, 15, 12, 42, 42, 16, 18, 18, 18, 18, 18,
        16, 16, 15, 18, 12, 22, 24, 22, 22, 24,
    ]
    wrapped_headers = {
        "TRIBUNAL",
        "NOMBRE",
        "NACIONALIDAD PAIS",
        "DERIVACION",
        "CURADOR",
    }

    def add_sheet(
        name: str,
        headers: list[str],
        rows: list[list[Any]],
        header_color: str,
        widths: list[int],
    ) -> None:
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_excel_value(value) for value in row])

        header_fill = PatternFill("solid", fgColor=header_color)
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 34
        last_column = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"

        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width

        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            required_lines = 1
            for column_index, cell in enumerate(row, start=1):
                header = normalize(headers[column_index - 1])
                wrap = header in wrapped_headers
                cell.alignment = Alignment(vertical="top", wrap_text=wrap)
                if wrap and cell.value not in {None, ""}:
                    required_lines = max(
                        required_lines,
                        math.ceil(len(str(cell.value)) / widths[column_index - 1]),
                    )
                if header.startswith("FEC ") and isinstance(cell.value, (date, datetime)):
                    cell.number_format = "dd-mm-yyyy"
            sheet.row_dimensions[row_number].height = min(max(20, required_lines * 15), 90)

    for status in ("Espera", "Cumplimiento"):
        add_sheet(status, HEADERS[status], records[status], "1F4E78", normal_widths)
    add_sheet("Duplicados", DUPLICATE_HEADERS, duplicate_rows, "9C6500", duplicate_widths)

    book.save(output_path)
    book.close()


def select_folder() -> Path | None:
    try:
        from selector_carpeta import seleccionar_carpeta_excel
    except ImportError as exc:
        raise RuntimeError("No se pudo cargar el selector de carpetas") from exc

    return seleccionar_carpeta_excel(
        "Selecciona la carpeta que contiene los archivos Excel"
    )


def show_message(title: str, message: str, error: bool = False) -> None:
    try:
        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        if error:
            messagebox.showerror(title, message)
        else:
            messagebox.showinfo(title, message)
        root.destroy()
    except Exception:
        pass


def run(folder: Path, dry_run: bool = False) -> Path | None:
    files = find_excels(folder)
    if not files:
        raise ValueError("No se encontraron archivos Excel en la carpeta o sus subcarpetas")

    extracted: list[ExtractedFile] = []
    errors: list[str] = []
    print(f"Archivos Excel encontrados: {len(files)}\n")
    for path in files:
        relative = path.relative_to(folder)
        try:
            item = extract_file(path)
            extracted.append(item)
            print(
                f"[OK] {relative} -> {item.status}: {len(item.records)} registro(s) "
                f"[hoja {item.sheet_name}, encabezado fila {item.header_row}]"
            )
        except Exception as exc:
            errors.append(f"{relative}: {exc}")
            print(f"[ERROR] {relative}: {exc}")

    if errors:
        detail = "\n".join(f"- {error}" for error in errors)
        raise ValueError(
            "No se generó el consolidado porque hay archivos que requieren revisión:\n"
            + detail
        )

    records: dict[str, list[list[Any]]] = {"Espera": [], "Cumplimiento": []}
    file_counts = {"Espera": 0, "Cumplimiento": 0}
    for item in extracted:
        records[item.status].extend(item.records)
        file_counts[item.status] += 1

    print("\nResumen:")
    for status in ("Espera", "Cumplimiento"):
        print(
            f"- {status}: {file_counts[status]} archivo(s), "
            f"{len(records[status])} registro(s)"
        )

    duplicate_rows, duplicate_groups = find_duplicate_records(records)
    print(
        f"- Duplicados: {duplicate_groups} coincidencia(s) por RUT + nombre, "
        f"{len(duplicate_rows)} registro(s) involucrado(s)"
    )

    if dry_run:
        print("\nSimulación terminada. No se creó el archivo de salida.")
        return None

    output_path = next_output_path(folder)
    create_output(output_path, records, duplicate_rows)
    print(f"\nArchivo creado: {output_path}")
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Consolida Excel de Espera y Cumplimiento en un solo archivo"
    )
    parser.add_argument("folder", nargs="?", help="Carpeta que se analizará")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Analiza los archivos sin crear el consolidado",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    selected_by_dialog = not args.folder
    try:
        folder = Path(args.folder).expanduser().resolve() if args.folder else select_folder()
        if folder is None:
            return 3
        if not folder.is_dir():
            raise ValueError(f"La ruta no es una carpeta válida: {folder}")
        output = run(folder, dry_run=args.dry_run)
        if selected_by_dialog and output is not None:
            show_message(
                "Consolidación terminada",
                f"Archivo creado correctamente:\n{output.name}",
            )
        return 0
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        if selected_by_dialog:
            show_message("No se creó el consolidado", str(exc), error=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
