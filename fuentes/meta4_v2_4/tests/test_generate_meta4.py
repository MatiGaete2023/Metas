"""Suite de regresión para generate_meta4.py v2.4.

Ejecutar desde la raíz del paquete:  python -m pytest tests -q
Todos los datos son ficticios. No usar archivos reales del Poder Judicial.

Cada bloque referencia el hallazgo de auditoría que protege (F-01, F-02, H-01,
H-02, H-03, M-02, etc.) para que el motivo de cada prueba sea rastreable.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
SCRIPT = BASE_DIR / "scripts" / "generate_meta4.py"
sys.path.insert(0, str(BASE_DIR / "scripts"))

import generate_meta4 as g  # noqa: E402

ESP_HDR = [
    "RIT", "TRIBUNAL", "RUT", "NOMBRE", "NACIONALIDAD / PAIS", "TIPO", "SEXO",
    "DERIVACIÓN", "CURADOR", "T ESPERA", "FEC. RESOLUCIÓN", "REVISIÓN", "DURACIÓN",
    "80 BIS", "FEC. NACIMIENTO", "EDAD", "FEC. ACT. F. INDIVIDUAL",
    "FEC. ACT. F. AMBULATORIA", "FEC. ACT. F. FAE/FAS", "FEC. OÍDO", "PROXS. AUDS.",
]

CUM_HDR = [
    "RIT", "TRIBUNAL", "RUT", "NOMBRE", "NACIONALIDAD / PAÍS", "TIPO", "SEXO",
    "DERIVACIÓN", "CURADOR", "FEC. RESOLUCIÓN", "FEC. INGRESO EFECTIVO",
    "FEC. EGRESO PROYECTADO", "DÍAS DE CUMPLIMIENTO", "DÍAS PARA EGRESAR", "80 BIS",
    "FEC. NACIMIENTO", "EDAD", "FEC. ACT. F. INDIVIDUAL", "FEC. ACT. F. AMBULATORIA",
    "FEC. ACT. F. FAE/FAS", "FEC. OÍDO",
]


def make_excel(
    path: Path, tribunal: str, estado: str, filas: list[dict], emision: str = "01/07/2026",
    when: dt.datetime | None = None,
) -> Path:
    headers = ESP_HDR if estado == "Espera" else CUM_HDR
    titulo = (
        "Informe de Ingresos en Lista de Espera"
        if estado == "Espera"
        else "Informe de Ingresos en Lista de Cumplimiento"
    )
    wb = Workbook()
    ws = wb.active
    ws["A1"] = f"Juzgado de Familia de {tribunal}"
    ws["A2"] = titulo
    ws["A3"] = f"Fecha Emisión   : {emision}"
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    for offset, fila in enumerate(filas):
        base = {
            "RIT": fila["rit"], "TRIBUNAL": tribunal, "RUT": "11111111-1",
            "NOMBRE": "FICTICIO NNA", "TIPO": "MP", "SEXO": "F",
            "DERIVACIÓN": fila["der"], "FEC. RESOLUCIÓN": dt.date(2026, 6, 15),
        }
        if estado == "Espera":
            base["T ESPERA"] = fila.get("tesp", 13)
        else:
            base["FEC. INGRESO EFECTIVO"] = dt.date(2026, 6, 20)
        for col, header in enumerate(headers, 1):
            ws.cell(6 + offset, col, base.get(header))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    if when:
        os.utime(path, (when.timestamp(), when.timestamp()))
    return path


def make_blob(path: Path, when: dt.datetime | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-ficticio" if path.suffix == ".pdf" else b"ficticio")
    if when:
        os.utime(path, (when.timestamp(), when.timestamp()))
    return path


def run(folder: Path, output: Path, *extra: str) -> dict:
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--folder", str(folder), "--output-dir", str(output),
         "--month", "2026-08", *extra],
        capture_output=True, text=True, check=True,
    )
    return json.loads(result.stdout)


def sheet_rows(path: Path, sheet: str) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is not None and str(row[0]).strip():
                rows.append(row)
        return rows
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Clasificación de tribunal por ruta (F-03, F-04)
# --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "ruta,esperado",
    [
        ("informes de la semana/espera fae x.pdf", None),
        ("INFORMES DE LA SEMANA/ESPERA FAE.pdf", None),
        ("ACTA DE LA REUNION.pdf", None),
        ("LA LISTA DE ESPERA.pdf", None),
        ("ESPERA FAE LA.pdf", "LOS ANGELES"),
        ("LA FAE CUMPLIMIENTO JULIO.xlsx", "LOS ANGELES"),
        ("CUMPLIMIENTO RES LA 2026.pdf", "LOS ANGELES"),
        ("espera fae la.pdf", "LOS ANGELES"),
        ("L.A. ESPERA FAE.pdf", "LOS ANGELES"),
        ("PROVINCIA CONCEPCION/cumplimiento residencia tome.pdf", "TOMÉ"),
        ("PROVINCIA DE ARAUCO - LEBU/espera fae.pdf", None),
        ("TOMEMOS NOTA.pdf", None),
        ("espera_fae_tome2026.xlsx", "TOMÉ"),
        ("LOS ANGELES/espera fae.pdf", "LOS ANGELES"),
        ("THNO/cumplimiento fae.pdf", "TALCAHUANO"),
        ("CCP/espera residencia.pdf", "CONCEPCIÓN"),
    ],
)
def test_tribunal_from_path(ruta, esperado):
    assert g.tribunal_from_path(ruta)[0] == esperado


def test_tribunal_name_celda():
    assert g.tribunal_name("Juzgado de Familia de Los Angeles") == "LOS ANGELES"
    assert g.tribunal_name("1° Juzgado de Familia de Concepción") == "CONCEPCIÓN"
    assert g.tribunal_name("Tribunal desconocido") is None


# --------------------------------------------------------------------------
# Estado y modalidad (F-06, F-07)
# --------------------------------------------------------------------------

def test_estado_ambiguo_no_clasifica():
    assert g.status_from_path("ESPERA Y CUMPLIMIENTO/x.pdf") is None
    assert g.status_from_path("cumplimiento/x.pdf") == "Cumplimiento"
    assert g.status_from_path("espera/x.pdf") == "Espera"


def test_modalidad_plural_y_ambigua():
    assert g.category_from_path("RESIDENCIAS/espera yumbel.pdf") == "Residencia"
    assert g.category_from_path("FAE/espera.pdf") == "FAE"
    assert g.category_from_path("FAE Y RESIDENCIA/espera.pdf") is None


# --------------------------------------------------------------------------
# Feriados (F-08)
# --------------------------------------------------------------------------

def test_feriado_solsticio_respaldo():
    assert dt.date(2026, 6, 21) in g.fallback_chile_holidays(2026)
    assert dt.date(2026, 6, 20) not in g.fallback_chile_holidays(2026)


def test_primer_dia_habil(tmp_path):
    assert g.first_business_day(2026, 1, BASE_DIR) == dt.date(2026, 1, 2)
    assert g.first_business_day(2026, 5, BASE_DIR) == dt.date(2026, 5, 4)
    assert g.first_business_day(2026, 8, BASE_DIR) == dt.date(2026, 8, 3)


# --------------------------------------------------------------------------
# F-01: sin duplicación de filas, huérfanas fuera de alcance
# --------------------------------------------------------------------------

def test_f01_sin_duplicacion_de_filas(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "FAE" / "lista espera coronel.xlsx", "CORONEL", "Espera",
        [{"rit": f"A-{i}-2026", "der": "DAM DIAGNOSTICO"} for i in range(1, 5)],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["verify_registros"]["rows"]["ESPERA"] == 0
    assert data["out_of_scope_rows"] == 4
    assert data["verification_problems"] == []
    assert sheet_rows(salida / "Registros.xlsx", "ESPERA") == []


def test_f01_mixto_conserva_solo_alcance(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "cumplimiento tome.xlsx", "TOMÉ", "Cumplimiento",
        [
            {"rit": "C-1-2026", "der": "FAE PRO"},
            {"rit": "C-2-2026", "der": "RESIDENCIA RLP"},
            {"rit": "C-3-2026", "der": "OPD SEGUIMIENTO"},
        ],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["verify_registros"]["rows"]["CUMPLIMIENTO"] == 2
    assert data["out_of_scope_rows"] == 1
    rits = {row[0] for row in sheet_rows(salida / "Registros.xlsx", "CUMPLIMIENTO")}
    assert rits == {"C-1-2026", "C-2-2026"}


# --------------------------------------------------------------------------
# F-02: un PDF nunca desplaza a un Excel con registros
# --------------------------------------------------------------------------

def test_f02_excel_vence_a_pdf_mas_reciente(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae lebu.xlsx", "LEBU", "Espera",
        [{"rit": f"B-{i}-2026", "der": "FAE PROGRAMA"} for i in range(1, 4)],
        emision="01/07/2026",
        when=dt.datetime(2026, 7, 1, 9, 0),  # explícitamente ANTES que el PDF
    )
    make_blob(folder / "espera fae lebu sin registros.pdf", dt.datetime(2026, 7, 30, 12, 0))
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["meta_summary"]["FAE Espera"] == 3
    assert data["verify_registros"]["rows"]["ESPERA"] == 3
    assert data["coverage"]["found"] == 1


def test_f02_error_si_seleccionado_pierde_registros(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "viejo" / "espera fae arauco.xlsx", "ARAUCO", "Espera",
        [{"rit": f"V-{i}-2026", "der": "FAE X"} for i in range(1, 4)], emision="01/06/2026",
    )
    make_excel(folder / "nuevo" / "espera fae arauco.xlsx", "ARAUCO", "Espera", [], emision="01/07/2026")
    salida = tmp_path / "out"
    run(folder, salida)
    reporte = next(salida.glob("Reporte de procesamiento*.txt")).read_text(encoding="utf-8-sig")
    assert "[ERROR]" in reporte
    assert "no aporta registros" in reporte


# --------------------------------------------------------------------------
# Continuidad, cobertura e idempotencia
# --------------------------------------------------------------------------

def test_faltantes_no_detienen_generacion(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "cumplimiento residencia laja.xlsx", "LAJA", "Cumplimiento",
        [{"rit": "G-1-2026", "der": "RESIDENCIA RLP"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 1
    assert data["coverage"]["missing"] == 67
    assert (salida / "Registros.xlsx").exists()
    assert list(salida.glob("META 4 - actualizado*.xlsx"))
    assert list(salida.glob("Reporte de procesamiento*.txt"))


def test_reejecucion_no_lee_sus_salidas(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae yumbel.xlsx", "YUMBEL", "Espera",
        [{"rit": "Y-1-2026", "der": "FAE UNO"}],
    )
    primera = run(folder, folder)
    segunda = run(folder, folder)
    assert primera["verify_registros"]["rows"]["ESPERA"] == 1
    assert segunda["verify_registros"]["rows"]["ESPERA"] == 1
    assert segunda["verification_problems"] == []


def test_pdf_incompleto_no_cubre_combinacion(tmp_path):
    folder = tmp_path / "in"
    make_blob(folder / "espera lebu.pdf")  # falta modalidad
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 0


def test_excel_corrupto_no_detiene(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir(parents=True)
    (folder / "roto.xlsx").write_bytes(b"esto no es un excel")
    make_excel(
        folder / "espera fae tome.xlsx", "TOMÉ", "Espera",
        [{"rit": "T-1-2026", "der": "FAE UNO"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["verify_registros"]["rows"]["ESPERA"] == 1
    reporte = next(salida.glob("Reporte de procesamiento*.txt")).read_text(encoding="utf-8-sig")
    assert "roto.xlsx" in reporte


def test_formato_t_espera_es_numero(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae laja.xlsx", "LAJA", "Espera",
        [{"rit": "L-1-2026", "der": "FAE UNO", "tesp": 13}],
    )
    salida = tmp_path / "out"
    run(folder, salida)
    meta = next(salida.glob("META 4 - actualizado*.xlsx"))
    wb = load_workbook(meta)
    try:
        ws = wb["FAE Espera"]
        celdas = [
            cell for row in ws.iter_rows(min_col=6, max_col=6)
            for cell in row if isinstance(cell.value, (int, float))
        ]
        assert celdas, "no se encontró la celda T ESPERA"
        assert all("YY" not in (cell.number_format or "").upper() for cell in celdas)
    finally:
        wb.close()


# --------------------------------------------------------------------------
# CLI (F-11)
# --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "argumentos,fragmento_esperado",
    [
        (["--month", "agosto"], "AAAA-MM"),
        (["--month", "2026-15"], "entre 01 y 12"),
        (["--month", "2026"], "AAAA-MM"),
        (["--emission-date", "32-13-2026"], "DD-MM-AAAA"),
    ],
)
def test_cli_argumento_invalido_sale_con_codigo_2(tmp_path, argumentos, fragmento_esperado):
    folder = tmp_path / "in"
    folder.mkdir(parents=True)
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--folder", str(folder), *argumentos],
        capture_output=True, text=True,
    )
    assert result.returncode == 2
    assert fragmento_esperado in result.stderr
    assert "Traceback" not in result.stderr


def test_cli_carpeta_inexistente_sale_con_codigo_1_sin_traceback(tmp_path):
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--folder", str(tmp_path / "no-existe")],
        capture_output=True, text=True,
    )
    assert result.returncode == 1
    assert "Traceback" not in result.stderr
    assert "No existe la carpeta" in result.stderr


def test_cli_ruta_no_es_carpeta_sale_con_codigo_1_sin_traceback(tmp_path):
    archivo = tmp_path / "no_es_carpeta.txt"
    archivo.write_text("x")
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--folder", str(archivo)],
        capture_output=True, text=True,
    )
    assert result.returncode == 1
    assert "Traceback" not in result.stderr


# --------------------------------------------------------------------------
# Selector de carpeta (v2.2)
# --------------------------------------------------------------------------

def test_pick_folder_cancelado_sale_con_codigo_3(monkeypatch, tmp_path, capsys):
    # No debe depender de un entorno gráfico real: se inyecta la cancelación
    # directamente en el proceso de prueba. Antes esta prueba invocaba
    # subprocess --pick-folder de verdad, lo que en Windows con tkinter
    # presente abre una ventana real y deja la prueba colgada (hallazgo M-05).
    monkeypatch.setattr(g, "select_source_folder", lambda: g.FolderSelection("CANCELADO"))
    monkeypatch.setattr(sys, "argv", ["generate_meta4.py", "--pick-folder"])
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit) as excinfo:
        g.main()
    assert excinfo.value.code == 3
    assert "cancelado" in capsys.readouterr().out.lower()


def test_pick_folder_no_disponible_sale_con_codigo_1(monkeypatch, tmp_path, capsys):
    monkeypatch.setattr(g, "select_source_folder", lambda: g.FolderSelection("NO_DISPONIBLE"))
    monkeypatch.setattr(sys, "argv", ["generate_meta4.py", "--pick-folder"])
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit) as excinfo:
        g.main()
    assert excinfo.value.code == 1
    assert "--folder" in capsys.readouterr().out


def test_folder_explicito_ignora_el_dialogo(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae laja.xlsx", "LAJA", "Espera",
        [{"rit": "P-1-2026", "der": "FAE UNO"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida, "--pick-folder")
    assert data["source_folder"] == str(folder.resolve())
    assert data["verify_registros"]["rows"]["ESPERA"] == 1


def test_sin_argumentos_usa_carpeta_actual(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae tome.xlsx", "TOMÉ", "Espera",
        [{"rit": "Q-1-2026", "der": "FAE UNO"}],
    )
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--month", "2026-08"],
        capture_output=True, text=True, cwd=str(folder), check=True,
    )
    data = json.loads(result.stdout)
    assert data["source_folder"] == str(folder.resolve())


def test_open_output_no_falla_fuera_de_windows(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae lebu.xlsx", "LEBU", "Espera",
        [{"rit": "R-1-2026", "der": "FAE UNO"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida, "--open-output")
    assert data["verify_registros"]["rows"]["ESPERA"] == 1


def test_selector_no_disponible_sin_dialogos(monkeypatch):
    monkeypatch.setattr(g, "_folder_dialog_tkinter", lambda: g.FolderSelection("NO_DISPONIBLE"))
    monkeypatch.setattr(g, "_folder_dialog_powershell", lambda: g.FolderSelection("NO_DISPONIBLE"))
    assert g.select_source_folder().status == "NO_DISPONIBLE"


def test_selector_usa_tkinter_primero(monkeypatch):
    monkeypatch.setattr(g, "_folder_dialog_tkinter", lambda: g.FolderSelection("SELECCIONADO", "/ruta/tk"))
    monkeypatch.setattr(g, "_folder_dialog_powershell", lambda: g.FolderSelection("SELECCIONADO", "/ruta/ps"))
    result = g.select_source_folder()
    assert result.status == "SELECCIONADO"
    assert result.path == "/ruta/tk"


def test_selector_no_disponible_cae_a_powershell(monkeypatch):
    monkeypatch.setattr(g, "_folder_dialog_tkinter", lambda: g.FolderSelection("NO_DISPONIBLE"))
    monkeypatch.setattr(g, "_folder_dialog_powershell", lambda: g.FolderSelection("SELECCIONADO", "/ruta/ps"))
    result = g.select_source_folder()
    assert result.status == "SELECCIONADO"
    assert result.path == "/ruta/ps"


def test_selector_cancelado_no_prueba_powershell(monkeypatch):
    # M-02: cancelar el primer diálogo no debe abrir un segundo. Antes, None
    # significaba a la vez "no disponible" y "cancelado", y ambos casos
    # activaban el respaldo de PowerShell.
    llamadas = []
    monkeypatch.setattr(g, "_folder_dialog_tkinter", lambda: g.FolderSelection("CANCELADO"))
    monkeypatch.setattr(
        g, "_folder_dialog_powershell", lambda: llamadas.append(1) or g.FolderSelection("SELECCIONADO", "/x")
    )
    result = g.select_source_folder()
    assert result.status == "CANCELADO"
    assert llamadas == []


# --------------------------------------------------------------------------
# H-01: el detector de duplicados no debe usar RIT+archivo como identidad
# --------------------------------------------------------------------------

def _build_registros_espera(path: Path, filas: list[list]) -> None:
    """Construye un Registros.xlsx mínimo (solo hoja ESPERA) para probar
    verify_registros_output de forma aislada, sin correr todo el pipeline."""
    headers = ESP_HDR + ["FECHA OBS.", "ARCHIVO ORIGEN", "RUTA ORIGEN"]
    wb = Workbook()
    ws = wb.active
    ws.title = "ESPERA"
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for row_idx, fila in enumerate(filas, 2):
        for col, value in enumerate(fila, 1):
            ws.cell(row_idx, col, value)
    wb.create_sheet("CUMPLIMIENTO")
    inv = wb.create_sheet("INVENTARIO_ARCHIVOS")
    inv.append(["N"])
    mat = wb.create_sheet("MATRIZ_COBERTURA")
    for _ in range(4):
        mat.append([""])
    wb.create_sheet("ERRORES_Y_ADVERTENCIAS")
    wb.save(path)


def _fila_base(rit, rut, nombre, derivacion, archivo):
    # RIT,TRIBUNAL,RUT,NOMBRE,NAC,TIPO,SEXO,DERIVACIÓN,CURADOR,T ESPERA,
    # FEC.RESOLUCIÓN,REVISIÓN,DURACIÓN,80BIS,FEC.NAC,EDAD,F.IND,F.AMB,F.FAE,
    # F.OÍDO,PROXS,FECHA OBS.,ARCHIVO ORIGEN,RUTA ORIGEN  (24 columnas)
    return [
        rit, "LAJA", rut, nombre, "CHILENA", "MP", "F", derivacion, "",
        13, dt.date(2026, 6, 15), "", "", "", dt.date(2015, 1, 1), 11,
        "", "", "", "", "", "", archivo, archivo,
    ]


def test_h01_hermanos_mismo_rit_no_es_falso_positivo(tmp_path):
    # Dos NNA distintos bajo la misma causa (mismo RIT): NO debe reportarse.
    path = tmp_path / "Registros.xlsx"
    _build_registros_espera(path, [
        _fila_base("F-100-2026", "11111111-1", "PERSONA FICTICIA UNO", "FAE UNO", "a.xlsx"),
        _fila_base("F-100-2026", "22222222-2", "PERSONA FICTICIA DOS", "FAE UNO", "a.xlsx"),
    ])
    result = g.verify_registros_output(path)
    assert result["problems"] == []
    assert result["rows"]["ESPERA"] == 2


def test_h01_fila_exacta_repetida_mismo_archivo_se_detecta(tmp_path):
    path = tmp_path / "Registros.xlsx"
    fila = _fila_base("F-200-2026", "33333333-3", "PERSONA FICTICIA TRES", "FAE UNO", "b.xlsx")
    _build_registros_espera(path, [fila, list(fila)])
    result = g.verify_registros_output(path)
    assert any("mismo archivo" in problema for problema in result["problems"])


def test_h01_fila_exacta_repetida_archivos_distintos_se_detecta(tmp_path):
    path = tmp_path / "Registros.xlsx"
    fila_a = _fila_base("F-300-2026", "44444444-4", "PERSONA FICTICIA CUATRO", "FAE UNO", "c.xlsx")
    fila_b = _fila_base("F-300-2026", "44444444-4", "PERSONA FICTICIA CUATRO", "FAE UNO", "d.xlsx")
    _build_registros_espera(path, [fila_a, fila_b])
    result = g.verify_registros_output(path)
    assert any("archivos distintos" in problema for problema in result["problems"])


def test_h01_muestra_julio_sin_falsos_positivos_sinteticos(tmp_path):
    # Reproduce la forma del hallazgo real: RIT repetido con fechas distintas
    # (misma persona, dos medidas) tampoco debe reportarse.
    path = tmp_path / "Registros.xlsx"
    f1 = _fila_base("F-400-2026", "55555555-5", "PERSONA FICTICIA CINCO", "FAE UNO", "e.xlsx")
    f2 = _fila_base("F-400-2026", "55555555-5", "PERSONA FICTICIA CINCO", "FAE UNO", "e.xlsx")
    f2[10] = dt.date(2026, 7, 1)  # FEC. RESOLUCIÓN distinta -> huella distinta
    _build_registros_espera(path, [f1, f2])
    result = g.verify_registros_output(path)
    assert result["problems"] == []


# --------------------------------------------------------------------------
# H-03: un Excel con más de un tribunal no debe contradecir la matriz
# --------------------------------------------------------------------------

def test_h03_excel_multitribunal_coherencia_matriz_meta(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae tome.xlsx", "TOMÉ", "Espera",
        [{"rit": "X-1-2026", "der": "FAE UNO"}],
    )
    wb = load_workbook(folder / "espera fae tome.xlsx")
    ws = wb.active
    ws.cell(7, 1, "X-2-2026")
    ws.cell(7, 2, "Juzgado de Familia de Lebu")
    ws.cell(7, 6, "MP")
    ws.cell(7, 7, "F")
    ws.cell(7, 8, "FAE UNO")
    ws.cell(7, 10, 5)
    ws.cell(7, 11, dt.date(2026, 6, 15))
    wb.save(folder / "espera fae tome.xlsx")

    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 2
    assert data["meta_summary"]["FAE Espera"] == 2
    assert data["verification_problems"] == []

    wb2 = load_workbook(next(salida.glob("Registros.xlsx")))
    matriz = wb2["MATRIZ_COBERTURA"]
    celdas = {}
    for row in matriz.iter_rows(min_row=5, max_row=21, values_only=True):
        if row[0] in ("TOMÉ", "LEBU"):
            celdas[row[0]] = row[1]
    assert "FALTANTE" not in (celdas.get("TOMÉ", ""))
    assert "FALTANTE" not in (celdas.get("LEBU", ""))


def test_h03_archivo_uniforme_sin_cambio_de_comportamiento(tmp_path):
    # Caso común: un solo tribunal en todo el archivo. Debe seguir dando un
    # único candidato, igual que antes de la corrección.
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae laja.xlsx", "LAJA", "Espera",
        [{"rit": f"U-{i}-2026", "der": "FAE UNO"} for i in range(1, 4)],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 1
    assert data["meta_summary"]["FAE Espera"] == 3
    assert data["verification_problems"] == []


# --------------------------------------------------------------------------
# H-02 opción A: Excel con datos vence a un PDF/imagen más reciente, pero
# la combinación queda visiblemente marcada para revisión humana.
# --------------------------------------------------------------------------

def test_h02a_excel_vence_y_queda_marcado_requiere_revision(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae lebu.xlsx", "LEBU", "Espera",
        [{"rit": f"B-{i}-2026", "der": "FAE PROGRAMA"} for i in range(1, 4)],
        emision="01/07/2026",
        when=dt.datetime(2026, 7, 1, 9, 0),
    )
    make_blob(folder / "espera fae lebu sin registros.pdf", dt.datetime(2026, 7, 30, 12, 0))
    salida = tmp_path / "out"
    run(folder, salida)

    wb = load_workbook(next(salida.glob("Registros.xlsx")))
    matriz = wb["MATRIZ_COBERTURA"]
    celda_lebu = next(row[1] for row in matriz.iter_rows(min_row=5, max_row=21, values_only=True) if row[0] == "LEBU")
    assert "REQUIERE REVISIÓN" in celda_lebu

    reporte = next(salida.glob("Reporte de procesamiento*.txt")).read_text(encoding="utf-8-sig")
    assert "[REQUIERE_REVISION]" in reporte
    assert "espera fae lebu.xlsx" in reporte
    assert "espera fae lebu sin registros.pdf" in reporte


# --------------------------------------------------------------------------
# M-05: casos de formato y de ruta que faltaban en la suite
# --------------------------------------------------------------------------

def test_xlsm_se_reconoce_y_procesa(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae laja.xlsm", "LAJA", "Espera",
        [{"rit": "M-1-2026", "der": "FAE UNO"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["verify_registros"]["rows"]["ESPERA"] == 1
    assert data["coverage"]["found"] == 1


def test_xls_legado_se_reconoce_y_procesa(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    folder = tmp_path / "in"
    folder.mkdir(parents=True)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Hoja1")
    ws.write(0, 0, "Juzgado de Familia de Laja")
    ws.write(1, 0, "Informe de Ingresos en Lista de Espera")
    ws.write(2, 0, "Fecha Emisión   : 01/07/2026")
    for col, header in enumerate(ESP_HDR, 0):
        ws.write(4, col, header)
    fila = {
        "RIT": "W-1-2026", "TRIBUNAL": "LAJA", "RUT": "11111111-1",
        "NOMBRE": "PERSONA FICTICIA", "TIPO": "MP", "SEXO": "F",
        "DERIVACIÓN": "FAE UNO", "T ESPERA": 13,
    }
    for col, header in enumerate(ESP_HDR, 0):
        ws.write(5, col, fila.get(header, ""))
    wb.save(str(folder / "espera fae laja.xls"))

    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 1
    assert data["verify_registros"]["rows"]["ESPERA"] == 1


def test_imagen_cubre_combinacion_sin_leer_contenido(tmp_path):
    folder = tmp_path / "in"
    make_blob(folder / "cumplimiento residencia laja.png")
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 1
    assert data["verify_registros"]["rows"]["CUMPLIMIENTO"] == 0
    wb = load_workbook(next(salida.glob("Registros.xlsx")))
    inv = wb["INVENTARIO_ARCHIVOS"]
    fila = next(r for r in inv.iter_rows(min_row=2, values_only=True) if r[1].endswith(".png"))
    assert fila[3] == "IMAGEN"
    assert fila[9] == "SIN REGISTROS"


def test_nombre_con_codificacion_dañada_usa_contenido_del_excel(tmp_path):
    # Simula un nombre de archivo corrompido por una extracción de ZIP mal
    # codificada (carácter de reemplazo Unicode). El tribunal debe salir del
    # CONTENIDO del Excel, no del nombre dañado -- ya que el nombre no permite
    # identificarlo de forma confiable.
    folder = tmp_path / "in"
    nombre_dañado = "espera fae \ufffdome.xlsx"
    make_excel(
        folder / nombre_dañado, "TOMÉ", "Espera",
        [{"rit": "D-1-2026", "der": "FAE UNO"}],
    )
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["verify_registros"]["rows"]["ESPERA"] == 1
    assert data["coverage"]["found"] == 1


def test_conflicto_carpeta_archivo_end_to_end(tmp_path):
    # M-01 de punta a punta: PDF cuyo nombre dice Cumplimiento pero cuya
    # carpeta dice Espera. El nombre del archivo (más específico) debe ganar.
    folder = tmp_path / "in"
    make_blob(folder / "ESPERA" / "cumplimiento residencia laja.pdf")
    salida = tmp_path / "out"
    data = run(folder, salida)
    assert data["coverage"]["found"] == 1
    wb = load_workbook(next(salida.glob("Registros.xlsx")))
    matriz = wb["MATRIZ_COBERTURA"]
    celda = next(
        r[4] for r in matriz.iter_rows(min_row=5, max_row=21, values_only=True) if r[0] == "LAJA"
    )
    assert celda != "FALTANTE"  # columna Residencia Cumplimiento cubierta
